"""Exportadores: Markdown -> md | docx | pdf, sin programas externos.

1. El Markdown se analiza con `mistune` y se reduce a una estructura intermedia simple (bloques + "runs"
   de texto con negrita/cursiva/código). Ambos formatos se generan desde esa misma estructura, así
   Word y PDF siempre muestran lo mismo.
2. Word: `python-docx` (MIT). Admite una plantilla de estilos (`reference_doc`).
3. PDF: `typst` (Apache-2.0), el compilador viene dentro del paquete de pip. El contenido se pasa como
   JSON a una plantilla fija (templates/document.typ), nunca como código: es seguro ante texto arbitrario.

Comentarios HTML `<!-- página N -->` (los pone el OCR) se convierten en saltos de página.
"""
from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path

import mistune

TEMPLATE = Path(__file__).parent / "templates" / "document.typ"
_PAGE_COMMENT = re.compile(r"^<!--\s*p[áa]gina\s+\d+\s*-->$", re.IGNORECASE)

_md = mistune.create_markdown(renderer="ast", plugins=["table", "strikethrough"])


def _run(text: str, bold: bool = False, italic: bool = False, code: bool = False) -> dict:
    return {"text": text, "bold": bold, "italic": italic, "code": code}


# ---------------------------------------------------------------- Markdown -> estructura intermedia


def _runs(children: list[dict] | None, bold: bool = False, italic: bool = False) -> list[dict]:
    out: list[dict] = []
    for node in children or []:
        t = node["type"]
        if t in ("text", "inline_html"):
            out.append(_run(node.get("raw", ""), bold, italic))
        elif t == "codespan":
            out.append(_run(node.get("raw", ""), bold, italic, code=True))
        elif t == "strong":
            out += _runs(node.get("children"), True, italic)
        elif t == "emphasis":
            out += _runs(node.get("children"), bold, True)
        elif t in ("link", "strikethrough"):
            out += _runs(node.get("children"), bold, italic)
        elif t == "image":
            alt = "".join(r["text"] for r in _runs(node.get("children")))
            out.append(_run(f"[imagen: {alt}]" if alt else "[imagen]", bold, True))
        elif t in ("softbreak", "linebreak"):
            out.append(_run(" ", bold, italic))
    # Unir runs contiguos con el mismo formato: menos nodos en Word y en Typst.
    merged: list[dict] = []
    for r in out:
        if merged and all(merged[-1][k] == r[k] for k in ("bold", "italic", "code")):
            merged[-1]["text"] += r["text"]
        else:
            merged.append(dict(r))
    return [r for r in merged if r["text"]]


def _list(node: dict) -> dict:
    items = []
    for item in node.get("children", []):
        runs: list[dict] = []
        children: list[dict] = []
        for child in item.get("children", []):
            if child["type"] in ("block_text", "paragraph"):
                if runs:
                    runs.append(_run(" "))
                runs += _runs(child.get("children"))
            else:
                children += _blocks([child])
        items.append({"runs": runs, "children": children})
    return {"type": "list", "ordered": bool(node.get("attrs", {}).get("ordered")), "items": items}


def _table(node: dict) -> dict | None:
    header: list[list[dict]] = []
    rows: list[list[list[dict]]] = []
    for part in node.get("children", []):
        if part["type"] == "table_head":
            header = [_runs(cell.get("children")) for cell in part.get("children", [])]
        elif part["type"] == "table_body":
            for row in part.get("children", []):
                rows.append([_runs(cell.get("children")) for cell in row.get("children", [])])
    if not header:
        return None
    width = len(header)
    rows = [(r + [[]] * width)[:width] for r in rows]  # filas irregulares -> mismo ancho que el encabezado
    return {"type": "table", "header": header, "rows": rows}


def _blocks(nodes: list[dict]) -> list[dict]:
    blocks: list[dict] = []
    for node in nodes:
        t = node["type"]
        if t == "heading":
            runs = _runs(node.get("children"))
            if runs:
                blocks.append({"type": "heading", "level": node.get("attrs", {}).get("level", 1), "runs": runs})
        elif t in ("paragraph", "block_text"):
            runs = _runs(node.get("children"))
            if runs:
                blocks.append({"type": "paragraph", "runs": runs})
        elif t == "list":
            blocks.append(_list(node))
        elif t == "table":
            table = _table(node)
            if table:
                blocks.append(table)
        elif t == "block_code":
            blocks.append({"type": "code", "text": node.get("raw", "").rstrip("\n")})
        elif t == "block_quote":
            runs: list[dict] = []
            for child in _blocks(node.get("children", [])):
                runs += child.get("runs", [])
            if runs:
                blocks.append({"type": "quote", "runs": runs})
        elif t == "block_html":
            if _PAGE_COMMENT.match(node.get("raw", "").strip()) and blocks:
                blocks.append({"type": "pagebreak"})
        elif t == "thematic_break":
            blocks.append({"type": "rule"})
    return blocks


_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


def _cells(line: str) -> int:
    return len(line.strip().strip("|").split("|"))


def normalize_tables(markdown: str) -> str:
    """Completa celdas faltantes en filas de tabla. GFM permite filas cortas, pero mistune descarta
    la tabla entera si una fila tiene menos celdas que el encabezado (pasa a menudo con OCR)."""
    lines = markdown.split("\n")
    out: list[str] = []
    width = 0
    for i, line in enumerate(lines):
        if _TABLE_RULE.match(line) and i > 0 and "|" in lines[i - 1]:
            width = _cells(line)
        elif width and (not line.strip() or "|" not in line):
            width = 0
        elif width:
            missing = width - _cells(line)
            if missing > 0:
                line = line.rstrip()
                line = (line if line.endswith("|") else line + " |") + " |" * missing
        out.append(line)
    return "\n".join(out)


def parse_blocks(markdown: str) -> list[dict]:
    return _blocks(_md(normalize_tables(markdown)))


def _title(blocks: list[dict], fallback: str) -> str:
    for b in blocks:
        if b["type"] == "heading":
            return "".join(r["text"] for r in b["runs"])
    return fallback


# ---------------------------------------------------------------- Word


def to_docx(blocks: list[dict], out: Path, title: str, reference_doc: Path | None = None) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document(str(reference_doc)) if reference_doc else Document()
    doc.core_properties.title = title

    def add_runs(paragraph, runs):
        for r in runs:
            run = paragraph.add_run(r["text"])
            run.bold = r["bold"] or None
            run.italic = r["italic"] or None
            if r["code"]:
                run.font.name = "Consolas"

    def add_list(block, depth):
        for i, item in enumerate(block["items"], 1):
            if block["ordered"]:
                # Numeración escrita a mano: el estilo "List Number" de Word continúa la cuenta entre listas.
                p = doc.add_paragraph(style="List Paragraph")
                p.paragraph_format.left_indent = Pt(18 * (depth + 1))
                p.add_run(f"{i}. ")
            else:
                p = doc.add_paragraph(style="List Bullet" if depth == 0 else f"List Bullet {min(depth + 1, 3)}")
            add_runs(p, item["runs"])
            for child in item["children"]:
                render(child, depth + 1)

    def render(b, depth=0):
        t = b["type"]
        if t == "heading":
            add_runs(doc.add_heading(level=min(b["level"], 9)), b["runs"])
        elif t == "paragraph":
            add_runs(doc.add_paragraph(), b["runs"])
        elif t == "list":
            add_list(b, depth)
        elif t == "table":
            table = doc.add_table(rows=1 + len(b["rows"]), cols=len(b["header"]))
            table.style = "Table Grid"
            for j, cell in enumerate(b["header"]):
                add_runs(table.rows[0].cells[j].paragraphs[0], [{**r, "bold": True} for r in cell])
            for i, row in enumerate(b["rows"], 1):
                for j, cell in enumerate(row):
                    add_runs(table.rows[i].cells[j].paragraphs[0], cell)
        elif t == "code":
            run = doc.add_paragraph().add_run(b["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif t == "quote":
            add_runs(doc.add_paragraph(style="Quote"), b["runs"])
        elif t == "pagebreak":
            doc.add_page_break()

    for block in blocks:
        render(block)
    doc.save(str(out))


# ---------------------------------------------------------------- PDF


def to_pdf(blocks: list[dict], out: Path, title: str) -> None:
    import typst

    payload = json.dumps({"title": title, "blocks": blocks}, ensure_ascii=False)
    out.write_bytes(typst.compile(str(TEMPLATE), sys_inputs={"doc": payload}))


# ---------------------------------------------------------------- API pública


def output_formats() -> dict[str, dict]:
    """Formatos de salida disponibles en esta instalación y, si no, qué falta instalar."""
    docx_ok = importlib.util.find_spec("docx") is not None
    pdf_ok = importlib.util.find_spec("typst") is not None
    return {
        "md": {"available": True, "reason": ""},
        "docx": {"available": docx_ok, "reason": "" if docx_ok else "Requiere python-docx"},
        "pdf": {"available": pdf_ok, "reason": "" if pdf_ok else "Requiere el paquete typst"},
    }


def export(markdown: str, stem: str, fmt: str, out_dir: Path, *, reference_doc: Path | None = None,
           timeout: int = 300) -> Path:
    md_path = out_dir / f"{stem}.md"
    md_path.write_text(markdown, encoding="utf-8")
    if fmt == "md":
        return md_path
    if fmt not in ("docx", "pdf"):
        raise ValueError(f"Formato de salida no soportado: {fmt}")

    blocks = parse_blocks(markdown)
    title = _title(blocks, stem)
    out = out_dir / f"{stem}.{fmt}"
    if fmt == "docx":
        to_docx(blocks, out, title, reference_doc)
    else:
        to_pdf(blocks, out, title)
    return out
