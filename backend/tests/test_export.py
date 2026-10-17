import io

from docx import Document
from pypdf import PdfReader

from app.modules.parse.export import export, normalize_tables, output_formats, parse_blocks

SAMPLE = """<!-- página 1 -->

# Informe de **Gestión**

Texto con *cursiva*, `código` y símbolos: #hashtag $100 // 50% <b>no html</b> @usuario [x].

- uno
- dos
  - dos punto uno

1. paso
2. paso

| Producto | Ventas | Variación |
|---|---:|---|
| Arroz | 1 200 | +5% |
| Azúcar | 830 |

<!-- página 2 -->

## Segunda página

Fin.
"""


def test_short_table_rows_are_padded():
    md = "| A | B | C |\n|---|---|---|\n| 1 | 2 |\n"
    assert normalize_tables(md).splitlines()[-1] == "| 1 | 2 | |"
    assert parse_blocks(md)[0]["type"] == "table"


def test_structure_is_preserved():
    blocks = parse_blocks(SAMPLE)
    types = [b["type"] for b in blocks]
    assert types == ["heading", "paragraph", "list", "list", "table", "pagebreak", "heading", "paragraph"]
    assert blocks[0]["runs"][1] == {"text": "Gestión", "bold": True, "italic": False, "code": False}
    assert blocks[2]["items"][1]["children"][0]["type"] == "list"  # lista anidada
    assert blocks[4]["rows"][1][2] == []  # celda faltante completada


def test_docx_export(tmp_path):
    out = export(SAMPLE, "informe", "docx", tmp_path)
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "Informe de Gestión" in texts
    assert any("#hashtag $100 // 50% <b>no html</b>" in t for t in texts)
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Producto", "Ventas", "Variación"]
    assert (tmp_path / "informe.md").is_file()  # el .md siempre se escribe


def test_pdf_export_is_safe_with_special_characters(tmp_path):
    out = export(SAMPLE, "informe", "pdf", tmp_path)
    reader = PdfReader(io.BytesIO(out.read_bytes()))
    assert len(reader.pages) == 2  # el comentario de página 2 es un salto de página
    text = reader.pages[0].extract_text()
    assert "#hashtag" in text and "$100" in text and "<b>no html</b>" in text


def test_all_formats_available_without_external_programs():
    assert all(f["available"] for f in output_formats().values())
